"""Extract hourly production and PTF from the Elestas workbook into tidy CSVs.

Usage: python scripts/extract_data.py [path/to/workbook.xlsx]
Outputs data/production_hourly.csv and data/ptf_hourly.csv.
"""
import sys
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/raw/Elestas_Saatlik_Uretimler_2025-2026.xlsx"

PLANTS = {
    "Araklı-1 HES": "HES", "Çataloluk HES": "HES", "Çobanlı HES": "HES",
    "Manahoz HES": "HES", "Murat HES": "HES", "Ortaçağ HES": "HES",
    "Polat HES": "HES", "Yaylabel HES": "HES", "Yazı HES": "HES",
    "GES Toplam": "GES", "Orhanlı RES": "RES", "Beylikova BES": "BES",
}


def read_plant_sheet(ws):
    """Plant sheets are blocks of: date row (with hour headers), 'MWh' row, blank row."""
    rows, day = [], None
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0] is None:
            continue
        if hasattr(r[0], "year"):
            day = r[0]
        elif r[0] == "MWh" and day is not None:
            for h in range(24):
                v = r[h + 1]
                rows.append((day + pd.Timedelta(hours=h), float(v) if isinstance(v, (int, float)) else None))
    return rows


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    frames = []
    for name, tech in PLANTS.items():
        df = pd.DataFrame(read_plant_sheet(wb[name]), columns=["datetime", "mwh"])
        df["plant"] = "GES (45 lisanssız)" if name == "GES Toplam" else name
        df["tech"] = tech
        frames.append(df)
    prod = pd.concat(frames, ignore_index=True)
    prod.to_csv(ROOT / "data/production_hourly.csv", index=False)

    ptf = []
    ws = wb["Analiz Veri 2025"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] is None:
            continue
        ptf.append((r[0] + pd.Timedelta(hours=int(r[1])), r[9], r[10], r[11]))
    ws = wb["Analiz Veri 2026 Fiyat"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] is None:
            continue
        ptf.append((r[0] + pd.Timedelta(hours=int(r[1])), r[10], r[11], r[12]))
    pd.DataFrame(ptf, columns=["datetime", "ptf_try", "ptf_usd", "ptf_eur"]).to_csv(
        ROOT / "data/ptf_hourly.csv", index=False)
    print(prod.groupby([prod.datetime.dt.year, "plant"]).mwh.agg(["sum", "count"]))
    print(len(ptf), "PTF hours")


if __name__ == "__main__":
    main()
